
import openpyxl
from openpyxl.utils import get_column_letter

FILE_PATH = 'Data/exogenous_data/DEA_Elec_Heat.xlsx'
# Use absolute path if necessary, but relative works if run from root. 
# The context says I am in C:\Users\borde\OneDrive\Bureau\model\EnergyScope_multi_cells
# I will use relative path.

SHEET_CONFIG = {
    '20 Onshore turbines': {
        'year_header_row': 2,
        'values': {
            'inv': 1.3,
            'fixed_om': 21000,
            'var_om': 3
        },
        'param_names': {
            'inv': 'Nominal investment (*total)',
            'fixed_om': 'Fixed O&M (*total)',
            'var_om': 'Variable O&M (*total)'
        }
    },
    '21 Offshore Wind AC Fixed': {
        'year_header_row': 2,
        'values': {
            'inv': 3.5,
            'fixed_om': 100000,
            'var_om': 5
        },
        'param_names': {
            'inv': 'Nominal investment (*total)',
            'fixed_om': 'Fixed O&M (*total)',
            'var_om': 'Variable O&M (*total)'
        }
    },
    '40 Comp. hp, airsource 10 MW': {
        'year_header_row': 2,
        'values': {
            'inv': 0.8,
            'fixed_om': 2000,
             # 'var_om' not specified for HP
        },
        'param_names': {
            'inv': 'Nominal investment (*total)',
            'fixed_om': 'Fixed O&M (*total)',
            'var_om': 'Variable O&M (*total)'
        }
    }
}

def find_param_row(ws, param_name, col_index=2):
    """Finds the row index where the cell in col_index contains param_name."""
    for row in ws.iter_rows(min_col=col_index, max_col=col_index):
        cell = row[0]
        if cell.value and isinstance(cell.value, str) and param_name in cell.value:
            return cell.row
    return None

def update_dea_file():
    print(f"Loading {FILE_PATH}...")
    try:
        wb = openpyxl.load_workbook(FILE_PATH)
    except FileNotFoundError:
        print(f"Error: File {FILE_PATH} not found.")
        return

    changes_made = False

    for sheet_name, config in SHEET_CONFIG.items():
        if sheet_name not in wb.sheetnames:
            print(f"Warning: Sheet '{sheet_name}' not found. Skipping.")
            continue
        
        ws = wb[sheet_name]
        print(f"\nProcessing sheet: {sheet_name}")

        header_row_idx = config['year_header_row']
        
        # Find 2015 and 2020 columns
        col_2015_idx = None
        col_target_copy_idx = None # The column to copy FROM (e.g. 2020)
        
        # Iterate over header row to find columns
        for cell in ws[header_row_idx]:
            if cell.value == 2015 or cell.value == "2015":
                col_2015_idx = cell.column
            
            # Look for 2020 or the first year > 2015 to use as source for copy
            if isinstance(cell.value, (int, float)):
                if cell.value == 2020:
                    col_target_copy_idx = cell.column
                elif col_target_copy_idx is None and cell.value > 2015:
                    col_target_copy_idx = cell.column # Fallback to next available year (e.g. 2025)

        target_col_idx = col_2015_idx

        if col_2015_idx:
            print(f"  - 2015 column found at column {get_column_letter(col_2015_idx)}.")
            # If it exists, we just update the specific values.
            # We assume other technical parameters are already there.
        else:
            # Insert new column
            if not col_target_copy_idx:
                print("  - Could not find a suitable source column (e.g. 2020) to copy from. Skipping insertion.")
                continue
            
            # We want to insert 'before' the source column generally, or chronologically.
            # Since 2015 < 2020, we insert before 2020.
            insert_idx = col_target_copy_idx
            
            print(f"  - Inserting 2015 column before column {get_column_letter(insert_idx)} (Year {ws.cell(header_row_idx, col_target_copy_idx).value}).")
            ws.insert_cols(insert_idx)
            
            target_col_idx = insert_idx
            # After insertion, the source column shifts right by 1
            source_col_idx = insert_idx + 1
            
            # Set header
            ws.cell(row=header_row_idx, column=target_col_idx).value = 2015
            
            # Copy all values from source column to new column
            # Note: We iterate all rows.
            for row in range(1, ws.max_row + 1):
                # Skip header row? No, header is already set, but we can overwrite or skip.
                if row == header_row_idx:
                    continue
                    
                source_cell = ws.cell(row=row, column=source_col_idx)
                target_cell = ws.cell(row=row, column=target_col_idx)
                target_cell.value = source_cell.value
                # Copy number format if needed
                if source_cell.number_format:
                    target_cell.number_format = source_cell.number_format

            changes_made = True

        # Now update specific parameters
        vals = config['values']
        params = config['param_names']
        
        # Investment
        if 'inv' in vals:
            row_idx = find_param_row(ws, params['inv'])
            if row_idx:
                print(f"  - Updating Investment (Row {row_idx}, Col {get_column_letter(target_col_idx)}) -> {vals['inv']}")
                ws.cell(row=row_idx, column=target_col_idx).value = vals['inv']
                changes_made = True
            else:
                print(f"  - Warning: Investment row '{params['inv']}' not found.")

        # Fixed O&M
        if 'fixed_om' in vals:
            row_idx = find_param_row(ws, params['fixed_om'])
            if row_idx:
                print(f"  - Updating Fixed O&M (Row {row_idx}, Col {get_column_letter(target_col_idx)}) -> {vals['fixed_om']}")
                ws.cell(row=row_idx, column=target_col_idx).value = vals['fixed_om']
                changes_made = True
            else:
                print(f"  - Warning: Fixed O&M row '{params['fixed_om']}' not found.")

        # Var O&M
        if 'var_om' in vals:
            row_idx = find_param_row(ws, params['var_om'])
            if row_idx:
                print(f"  - Updating Var O&M (Row {row_idx}, Col {get_column_letter(target_col_idx)}) -> {vals['var_om']}")
                ws.cell(row=row_idx, column=target_col_idx).value = vals['var_om']
                changes_made = True
            else:
                print(f"  - Warning: Var O&M row '{params['var_om']}' not found.")

    print("\nSaving file...")
    wb.save(FILE_PATH)
    print("Updates completed successfully.")

if __name__ == "__main__":
    update_dea_file()
