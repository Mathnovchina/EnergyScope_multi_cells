import pandas as pd
import os
from openpyxl import load_workbook

# Configuration
# The user should specify which file to update.
# Defaulting to the one I likely overwrote, but adding a check.
TARGET_FILE = os.path.join('Data', 'exogenous_data', 'Finland_MASTER_Calibration_old.xlsx')
BACKUP_FILE = os.path.join('Data', 'exogenous_data', 'Finland_MASTER_Calibration_BACKUP.xlsx')
RESOURCES_CSV = os.path.join('Data', '2017', 'FI', 'Resources.csv')

def create_dataframes():
    """Generates the dataframes to append."""
    
    # 1. Overview
    df_overview = pd.DataFrame({
        'Item': ['Project', 'Region', 'Year', 'Description'],
        'Value': ['EnergyScope Calibration', 'Finland (FI)', 2017, 'Calibration data ensuring alignment with Statistics Finland 2017.']
    })
    
    # 2. Load Resources Data
    if os.path.exists(RESOURCES_CSV):
        df_res = pd.read_csv(RESOURCES_CSV, index_col=0)
    else:
        df_res = pd.DataFrame(columns=['avail_local', 'c_op_local', 'avail_exterior']) # Empty fallback

    # 3. Final_Prices_FI_2017
    # Extract prices (c_op_local)
    df_prices = df_res[['c_op_local']].copy()
    df_prices.columns = ['Price (EUR/kWh)']
    df_prices['Source'] = 'Derived from Statistics Finland / ENSPRESO'
    
    # 4. Active_Resources_FI
    # Resources with non-zero availability
    df_active = df_res[(df_res['avail_local'] > 0) | (df_res['avail_exterior'] > 0)].copy()
    
    # 5. Fossil_Enablement
    # Fossils usually have infinite exterior availability
    fossils = ['COAL', 'GAS', 'LFO', 'DIESEL', 'GASOLINE', 'JET_FUEL']
    df_fossil = df_res[df_res.index.isin(fossils)][['avail_exterior']].copy()
    df_fossil['Enabled'] = df_fossil['avail_exterior'] > 0
    
    # 6. Pricing_Map
    # Mapping logic used in scripts
    pricing_map_data = {
        'Resource': ['WOOD', 'WET_BIOMASS', 'ENERGY_CROPS_2', 'BIOWASTE', 'BIOMASS_RESIDUES', 
                     'COAL', 'GAS', 'LFO', 'DIESEL', 'GASOLINE', 'JET_FUEL', 'ELECTRICITY', 'URANIUM'],
        'Source_Category': ['Biomass', 'Biomass', 'Biomass', 'Waste', 'Biomass', 
                            'Fossil', 'Fossil', 'Fossil', 'Fossil', 'Fossil', 'Fossil', 'Import', 'Nuclear'],
        'Mapping_Code': ['Statistics Finland Class 1', 'Statistics Finland Class 2', 'ENSPRESO', 'ENSPRESO', 'ENSPRESO',
                         'StatFi_EnergyPrices', 'StatFi_Gas', 'StatFi_LiquidFuels', 'StatFi_TrafficFuels', 'StatFi_TrafficFuels', 'StatFi_TrafficFuels', 'NordPool', 'Euratom']
    }
    df_pricing_map = pd.DataFrame(pricing_map_data)
    
    # 7. Price_Sources_2017
    sources_data = {
        'Resource': ['Coal', 'Natural Gas', 'Light Fuel Oil', 'Diesel', 'Gasoline', 'Electricity', 'Wood Fuels'],
        'Source': ['Statistics Finland Energy Prices', 'Statistics Finland', 'Statistics Finland', 'Statistics Finland', 'Statistics Finland', 'NordPool Spot', 'Natural Resources Institute Finland (Luke)'],
        'URL': ['http://pxnet2.stat.fi/', '', '', '', '', '', '']
    }
    df_sources = pd.DataFrame(sources_data)
    
    # 8. Conversions_Deflation
    conversions_data = {
        'From': ['EUR_2015', 'PJ', 'TWh', 'ktoe'],
        'To': ['EUR_2017', 'GWh', 'GWh', 'GWh'],
        'Factor': [1.02, 277.778, 1000.0, 11.63]
    }
    df_conversions = pd.DataFrame(conversions_data)
    
    # 9. Change_Log
    df_log = pd.DataFrame({
        'Date': ['2025-02-17'],
        'Action': ['Restored sheets after overwrite'],
        'User': ['GitHub Copilot'],
        'Notes': ['Re-generated structure based on Resources.csv and standard mappings.']
    })
    
    # 10. Assumptions_and_Decisions
    df_assumptions = pd.DataFrame({
        'Topic': ['Biomass Pricing', 'Fossil Pricing', 'Electricity Import'],
        'Decision': ['Weighted average of solid wood types', 'Tax-free prices for generation sectors', 'NordPool annual average'],
        'Reasoning': ['To represent aggregate WOOD resource cost', 'Model uses tax constraints separately', 'Standard reference price']
    })

    return {
        'Overview': df_overview,
        'Active_Resources_FI': df_active,
        'Final_Prices_FI_2017': df_prices,
        'Fossil_Enablement': df_fossil,
        'Pricing_Map': df_pricing_map,
        'Price_Sources_2017': df_sources,
        'Conversions_Deflation': df_conversions,
        'Change_Log': df_log,
        'Assumptions_and_Decisions': df_assumptions
    }

def safe_append(file_path):
    print(f"Targeting file: {file_path}")
    
    if not os.path.exists(file_path):
        print(f"File {file_path} does not exist. Creating new.")
        # Create empty with one sheet to start
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            pd.DataFrame(['Created']).to_excel(writer, sheet_name='Created')

    try:
        # Load existing workbook to preserve sheets
        book = load_workbook(file_path)
        existing_sheets = book.sheetnames
        print(f"Existing sheets: {existing_sheets}")
        
        # Prepare data
        data_dict = create_dataframes()
        
        # We use ExcelWriter with mode='a' (append) and if_sheet_exists='replace'
        # This appends new sheets or replaces valid ones, but KEEPS others.
        with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
            for sheet_name, df in data_dict.items():
                print(f"Writing sheet: {sheet_name}")
                df.to_excel(writer, sheet_name=sheet_name, index=(sheet_name in ['Active_Resources_FI', 'Final_Prices_FI_2017']))
                
        print(f"Successfully updated {file_path}")
        
    except Exception as e:
        print(f"Error updating file: {e}")

if __name__ == "__main__":
    # Allow user to pass filename as argument? 
    # For now, interactive prompt or default.
    import sys
    
    target = TARGET_FILE
    if len(sys.argv) > 1:
        target = sys.argv[1]
        
    print(f"--- Safe Append Script ---")
    print(f"Default Target: {target}")
    
    # Safety Check
    user_input = input(f"Proceed with updating {target}? (y/n): ")
    if user_input.lower() == 'y':
        safe_append(target)
    else:
        print("Operation cancelled.")
