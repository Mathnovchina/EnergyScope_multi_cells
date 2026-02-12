import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
import os

file_path = r'c:\Users\borde\OneDrive\Bureau\model\EnergyScope_multi_cells\Data\exogenous_data\Finland_MASTER_Calibration.xlsx'
sheet_name = '2_Master_Log'

def log_changes():
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    # Data to append
    new_rows = [
        {
            'Category': 'Data Update',
            'Parameter': 'Technology Costs (Inv, Fixed O&M)',
            'Value': 'Updated to DEA 2023 Catalogue (Interpolated to 2017)',
            'Source/Methodology': 'scripts/fix_technologies_errors.py processing DEA_Elec_Heat.xlsx. Mapped correctly (c_p preserved).',
            'Date': datetime.now().strftime('%Y-%m-%d')
        },
        {
            'Category': 'Data Update',
            'Parameter': 'Resources Import Prices',
            'Value': 'Updated to 2017 Historical Averages',
            'Source/Methodology': 'Manual/Script update to Resources.csv based on NordPool/StatFi data.',
            'Date': datetime.now().strftime('%Y-%m-%d')
        },
         {
            'Category': 'Data Update',
            'Parameter': 'Network Exchanges',
            'Value': 'FI-EE Gas Capacity set to 0 (No Balticconnector in 2017)',
            'Source/Methodology': 'Update to Network_exchanges.csv',
            'Date': datetime.now().strftime('%Y-%m-%d')
        },
        {
            'Category': 'Model Run',
            'Parameter': 'Execution status',
            'Value': 'Solved / Validated',
            'Source/Methodology': 'scripts/run_finland.py -> case_studies/FI/ref_2017_finland',
            'Date': datetime.now().strftime('%Y-%m-%d')
        }
    ]

    df_new = pd.DataFrame(new_rows)
    
    # Append using openpyxl directly to preserve file
    print(f"Appending {len(new_rows)} rows to {sheet_name} in {file_path}")
    
    try:
        book = load_workbook(file_path)
        if sheet_name not in book.sheetnames:
            print(f"Sheet {sheet_name} not found. Creating it.")
            ws = book.create_sheet(sheet_name)
            # Add header if new
            ws.append(df_new.columns.tolist())
        else:
            ws = book[sheet_name]
            
        # Append rows
        for index, row in df_new.iterrows():
            ws.append(row.tolist())
            
        book.save(file_path)
        print("Success.")
        
    except Exception as e:
        print(f"Error appending to Excel: {e}")

if __name__ == "__main__":
    log_changes()
