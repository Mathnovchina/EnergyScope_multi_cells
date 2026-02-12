
import openpyxl
from datetime import datetime

file_path = r'c:\Users\borde\OneDrive\Bureau\model\EnergyScope_multi_cells\Data\exogenous_data\Finland_Calibration_MASTER.xlsx'

def main():
    print(f"Updating {file_path}")
    today = datetime.now().strftime("%Y-%m-%d")
    
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return

    # 1. Critical Changes
    sheet_name = '3_Critical_Changes'
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        data = [
            ['High', 'WIND_ONSHORE 2017', '2035 constraint (f_min > 3GW) infeasible for 2017 (<2GW)', 'Update f_min using 2017 specific data', 'Done', 'Critical - Feasibility'],
            ['High', 'NUCLEAR 2035/2050', 'Forced to 0 GW in default config', 'Updated CSVs with realistic capacity (4.36 GW)', 'Done', 'Critical - Base Load']
        ]
        for row in data:
            ws.append(row)
        print(f"Appended to {sheet_name}")

    # 2. Next Steps
    sheet_name = '11_Next_Steps'
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        data = [
            ['Calibration', 1, 'Initial 2017 Model Run', 'Ready', '1 day', 'Data structure complete'],
            ['Post-Processing', 2, 'Validate 2017 Output vs Stats', 'Pending', '1 day', 'Check production mix']
        ]
        for row in data:
            ws.append(row)
        print(f"Appended to {sheet_name}")

    # 3. Issues Decisions
    sheet_name = '12_Issues_Decisions'
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        data = [
            [today, '2017 demand data source', 'Use Energy in Finland 2022 (via Calibration file)', 'Validated source', 'Closed'],
            [today, 'Cost Learning Curves', 'Use static 2035 costs for now', 'No learning script in repo. Impact considered low for 2017 calib.', 'Open']
        ]
        for row in data:
            ws.append(row)
        print(f"Appended to {sheet_name}")

    try:
        wb.save(file_path)
        print("Workbook saved successfully.")
    except Exception as e:
        print(f"Error saving workbook: {e}")

if __name__ == "__main__":
    main()
