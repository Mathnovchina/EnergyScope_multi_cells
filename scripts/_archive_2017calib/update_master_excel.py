"""
Update Finland_MASTER_Calibration_old_UPDATED.xlsx with:
  - 20_V37_Tech_Constraints  : All technology fmin/fmax/fperc after applying v37 patch chain
  - 21_V37_Resource_Inputs   : All resource avail + prices after applying v37 patch chain, with sources
  - Updated 12_Validation_2017 with v37 scorecard
  - Updated 18_Change_Log
"""
import csv, os, copy
from datetime import date
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL = r"Data\exogenous_data\Finland_MASTER_Calibration_old_UPDATED.xlsx"

# ── Patch chain (order matters) ──────────────────────────────────────────────
PATCH_DIR = "calibration/patches"
PATCH_CHAIN = [
    "restore_v9_baseline.csv",
    "v10_delta.csv",
    "v12_disable_future_techs.csv",
    "v16_cap_coal.csv",
    "v17_nuclear_fmin.csv",
    "v20_disable_biomass_hvc.csv",
    "v21_wind_cap.csv",
    "v22_gas_cap.csv",
    "v23_ind_cogen_wood.csv",
    "v24_hydro_cap.csv",
    "v25_ind_cogen_wood2.csv",
    "v26_gas_rebalance.csv",
    "v27_ind_cogen_wood3.csv",
    "v28_coal_cap.csv",
    "v29_nuclear_solar.csv",
    "v33_oil_biomass3.csv",
    "v36_grid_losses_recalib.csv",
    "v37_solar.csv",
]

# ── Read patches ─────────────────────────────────────────────────────────────
def read_patches():
    """Return dict of (file, param, tech) -> (value, patch_source)."""
    result = OrderedDict()
    for pf in PATCH_CHAIN:
        path = os.path.join(PATCH_DIR, pf)
        with open(path, newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                key = (row["file"].strip(), row["parameter"].strip(), row["technology_or_resource"].strip())
                result[key] = (float(row["value"].strip()), pf)
    return result

# ── Classify patches ─────────────────────────────────────────────────────────
def classify_patches(patches):
    tech_rows = {}   # tech_name -> {param: (value, source)}
    res_rows  = {}   # res_name  -> {param: (value, source)}
    for (fpath, param, name), (val, src) in patches.items():
        if "Technologies" in fpath:
            tech_rows.setdefault(name, {})[param] = (val, src)
        elif "Resources" in fpath:
            res_rows.setdefault(name, {})[param] = (val, src)
    return tech_rows, res_rows

# ── Source annotations for resources ─────────────────────────────────────────
RESOURCE_SOURCES = {
    "WOOD":            ("ENSPRESO + Finnish Forest Research Institute (Luke)", "Local forestry residues + energy wood potential for 2020"),
    "WET_BIOMASS":     ("ENSPRESO / local estimates", "Agricultural wet biomass, biogas feedstock"),
    "ENERGY_CROPS_2":  ("ENSPRESO / local estimates", "Energy crops potential for Finland"),
    "BIOWASTE":        ("ENSPRESO / local estimates", "Municipal + industrial biowaste"),
    "BIOMASS_RESIDUES":("ENSPRESO / local estimates", "Agricultural residues"),
    "WASTE":           ("Statistics Finland + Eurostat", "Municipal solid waste energy content"),
    "ELECTRICITY":     ("Market import — uncapped", "Nord Pool interconnections (SE, EE, RU); import_capacity=4.5 GW from Misc.json"),
    "GAS":             ("Statistics Finland Energy Balance 2017", "Natural gas supply capped at 20 TWh to match 2017 consumption (~20 TWh); applied via v22_gas_cap.csv"),
    "COAL":            ("Statistics Finland Energy Balance 2017", "Coal supply capped at 35 TWh to match 2017 consumption (~35 TWh total coal+peat); applied via v16_cap_coal.csv"),
    "URANIUM":         ("IAEA PRIS / TVO+Fortum", "Nuclear fuel — uncapped, cost=0.0093 MEUR/GWh from DEA/JRC"),
    "LFO":             ("Market import — uncapped", "Light fuel oil, price=0.0521 MEUR/GWh from Statistics Finland fuel prices"),
    "DIESEL":          ("Market import — uncapped", "Transport diesel, price=0.0543 MEUR/GWh from Statistics Finland"),
    "GASOLINE":        ("Market import — uncapped", "Transport gasoline, price=0.0588 MEUR/GWh from Statistics Finland"),
    "JET_FUEL":        ("Market import — uncapped", "Aviation fuel, price=0.0359 MEUR/GWh (wholesale, not pump-price)"),
    "H2":              ("Not available in 2017", "H2 import available for future scenarios only"),
    "RES_HYDRO":       ("Statistics Finland + Fingrid", "Hydro resource capped at 14,600 GWh to match 2017 production; applied via v24_hydro_cap.csv"),
}

# ── Technology calibration rationale ─────────────────────────────────────────
TECH_RATIONALE = {
    "NUCLEAR":           "2.7–2.8 GW matches 2017 fleet: Loviisa 1&2 (1.01 GW) + Olkiluoto 1&2 (1.76 GW); OL3 not yet online",
    "PV_ROOFTOP":        "f_min=0.055 GW: ~55 MW installed in 2017 (Finnish Energy Authority); raised from 0.05 in v37 to match ELEC_SOLAR output; f_max from base",
    "PV_UTILITY":        "f_max=1.0 GW from base; negligible utility PV in 2017",
    "WIND_ONSHORE":      "1.4–1.59 GW matches 2017 installed capacity (~2.0 GW nameplate, ~1.5 GW effective); Statistics Finland + Finnish Wind Power Association",
    "WIND_OFFSHORE":     "f_max=0 — no offshore wind in Finland in 2017",
    "HYDRO_DAM":         "1.1–1.3 GW from base; Kemijoki etc.",
    "HYDRO_RIVER":       "1.9–2.1 GW from base; run-of-river fleet",
    "COAL_US":           "f_min=3.0, f_max=4.5 GW; fmax_perc=0.047 — coal fleet still active in 2017, capped to match Statistics Finland electricity output (~3.3 TWh)",
    "CCGT":              "f_max=1.5 GW from base; gas CCGT plants, marginal role in 2017",
    "DHN_COGEN_GAS":     "f_min=1.42 GW matches installed gas-CHP fleet for DHN (Helsinki, Vantaa) — raised from 1.29 in v36 grid-loss recalibration; fmax_perc=0.2 from base",
    "DHN_COGEN_WOOD":    "fmin_perc=0.2 — significant bio-CHP fleet (Jyväskylä, Kuopio, etc.); Statistics Finland CHP data",
    "DHN_COGEN_COAL":    "fmin_perc=0.25 — coal-CHP still operating in Helsinki (Salmisaari, Hanasaari), Naantali",
    "DHN_BOILER_OIL":    "fmin_perc=0.05, fmax_perc=0.1 — oil backup boilers for DHN peak demand",
    "DHN_BOILER_WOOD":   "fmin_perc=0.15 — wood-fired boilers for DHN, common in Finnish municipalities",
    "IND_COGEN_WOOD":    "fmin_perc=0.55 — Finnish pulp & paper industry operates large wood-CHP fleet; Statistics Finland industrial CHP data",
    "IND_BOILER_WOOD":   "fmin_perc=0.2 — industrial wood boilers dominant in Finnish industry; forestry sector tradition",
    "IND_BOILER_OIL":    "f_min=0.2, fmin_perc=0.05 — industrial oil boilers for peak/backup",
    "IND_BOILER_COAL":   "fmin_perc=0.05 — some industrial coal boilers remain in 2017",
    "IND_BOILER_GAS":    "fmax_perc=0.2 — gas limited in industrial heat; gas mainly for CHP and DEC boilers",
    "DEC_BOILER_GAS":    "fmax_perc=0.18 — gas boilers in decentralized heating limited to ~18% share; v26_gas_rebalance",
    "DEC_BOILER_OIL":    "fmin_perc=0.10 — oil heating remains significant in Finnish households, especially rural areas",
    "DEC_BOILER_WOOD":   "fmax_perc=0.25, fmin_perc=0.20 — wood heating capped to balance biomass use vs. oil floor; fmin_perc=0.20 added in v36 to prevent HP displacement after grid-loss correction",
    "DEC_DIRECT_ELEC":   "fmax_perc=0.6 — electric heating common in Finland (sähkölämmitys), ~30-40% market share",
    "DEC_SOLAR":         "f_max=0.1 GW — negligible solar thermal in 2017 Finland",
    "DEC_THHP_GAS":      "fmax_perc=0 — gas heat pumps not deployed in Finland in 2017",
    "GEOTHERMAL":        "f_max=0.3 GW from base; no significant geothermal electricity in Finland",
    "SYN_METHANATION":   "f_max=0 — not available in 2017, disabled via v12_disable_future_techs",
    "BIOMASS_TO_HVC":    "f_max=0 — biomass-to-chemicals pathway not available in 2017; disabled via v20",
    "CAR_GASOLINE":      "fmin_perc=0.58, fmax_perc=0.65 — Statistics Finland vehicle fleet: ~60% gasoline",
    "CAR_DIESEL":        "fmin_perc=0.38, fmax_perc=0.40 — Statistics Finland vehicle fleet: ~38% diesel",
    "CAR_BEV":           "fmax_perc=0.01 — BEVs negligible in 2017 (<1%)",
    "CAR_PHEV":          "fmax_perc=0.01 — PHEVs negligible in 2017",
    "CAR_HEV":           "fmax_perc=0.05 — small hybrid share",
    "CAR_FUEL_CELL":     "fmax_perc=0 — no FCEV fleet in 2017",
    "TRUCK_DIESEL":      "fmin_perc=0.95 — almost all trucks diesel in 2017",
    "TRUCK_NG":          "fmax_perc=0.05 — tiny CNG truck share",
    "BUS_COACH_DIESEL":  "fmin_perc=0.95 — buses overwhelmingly diesel in 2017",
    "CARGO_LFO":         "f_min=15 GW, fmin_perc=0.95 — international shipping LFO dominant",
    "CARGO_LNG":         "fmax_perc=0.05 — LNG shipping marginal",
    "BOAT_FREIGHT_DIESEL":"fmin_perc=0.95 — inland/coastal freight diesel dominant",
    "BOAT_FREIGHT_NG":   "fmax_perc=0.05 — LNG inland shipping marginal",
    "TRAIN_PUB":         "fmax_perc=0.5 — rail in public transport capped per modal data",
    "TRAMWAY_TROLLEY":   "fmax_perc=0.3 — Helsinki tram system share of public transport",
}

# ── Build effective Tech constraints after all patches ───────────────────────
def build_effective_tech(tech_patches):
    """Group by technology, show final effective values + patch source."""
    rows = []
    for tech, params in sorted(tech_patches.items()):
        f_min = params.get("f_min", (None, ""))
        f_max = params.get("f_max", (None, ""))
        fmin_p = params.get("fmin_perc", (None, ""))
        fmax_p = params.get("fmax_perc", (None, ""))
        rationale = TECH_RATIONALE.get(tech, "")
        rows.append({
            "Technology": tech,
            "f_min": f_min[0],
            "f_min_source": f_min[1] if f_min[0] is not None else "",
            "f_max": f_max[0],
            "f_max_source": f_max[1] if f_max[0] is not None else "",
            "fmin_perc": fmin_p[0],
            "fmin_perc_source": fmin_p[1] if fmin_p[0] is not None else "",
            "fmax_perc": fmax_p[0],
            "fmax_perc_source": fmax_p[1] if fmax_p[0] is not None else "",
            "Rationale": rationale,
        })
    return rows

# ── Build effective Resource constraints ─────────────────────────────────────
def build_effective_res(res_patches):
    rows = []
    for res, params in sorted(res_patches.items()):
        avail_local = params.get("avail_local", (None, ""))
        avail_ext   = params.get("avail_exterior", (None, ""))
        c_op_local  = params.get("c_op_local", (None, ""))
        source_info = RESOURCE_SOURCES.get(res, ("", ""))
        rows.append({
            "Resource": res,
            "avail_local": avail_local[0],
            "avail_local_source": avail_local[1] if avail_local[0] is not None else "",
            "avail_exterior": avail_ext[0],
            "avail_exterior_source": avail_ext[1] if avail_ext[0] is not None else "",
            "c_op_local": c_op_local[0],
            "c_op_local_source": c_op_local[1] if c_op_local[0] is not None else "",
            "Source": source_info[0],
            "Notes": source_info[1],
        })
    return rows

# ── Style helpers ────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
WRAP = Alignment(wrap_text=True, vertical="top")

def write_header(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w

def write_row(ws, row, values, bold=False):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.border = THIN_BORDER
        cell.alignment = WRAP
        if bold:
            cell.font = Font(bold=True)

# ── MAIN ─────────────────────────────────────────────────────────────────────
def main():
    patches = read_patches()
    tech_patches, res_patches = classify_patches(patches)
    tech_rows = build_effective_tech(tech_patches)
    res_rows  = build_effective_res(res_patches)

    wb = openpyxl.load_workbook(EXCEL)

    # ── Sheet 20: Technology Constraints ──────────────────────────────────
    if "20_V33_Tech_Constraints" in wb.sheetnames:
        del wb["20_V33_Tech_Constraints"]
    if "20_V37_Tech_Constraints" in wb.sheetnames:
        del wb["20_V37_Tech_Constraints"]
    ws_tech = wb.create_sheet("20_V37_Tech_Constraints")

    # Title
    ws_tech.cell(1, 1, "Technology Constraints — v37 Calibration (Finland 2017)")
    ws_tech.cell(1, 1).font = Font(bold=True, size=14)
    ws_tech.merge_cells("A1:J1")
    ws_tech.cell(2, 1, f"Patch chain: {' → '.join(p.replace('.csv','') for p in PATCH_CHAIN)}")
    ws_tech.cell(2, 1).font = Font(italic=True, size=9)
    ws_tech.merge_cells("A2:J2")
    ws_tech.cell(3, 1, f"Generated: {date.today().isoformat()}")
    ws_tech.cell(3, 1).font = Font(italic=True, size=9, color="888888")

    headers = [
        "Technology", "f_min (GW)", "Source (f_min)", "f_max (GW)", "Source (f_max)",
        "fmin_perc", "Source (fmin_perc)", "fmax_perc", "Source (fmax_perc)", "Rationale / Source"
    ]
    widths = [24, 12, 22, 14, 22, 12, 22, 12, 22, 55]
    write_header(ws_tech, 5, headers, widths)

    # Separate into categories
    energy_techs = []
    transport_techs = []
    infra_techs = []
    synfuel_techs = []
    storage_techs = []
    other_techs = []

    TRANSPORT_PREFIXES = ("CAR_", "TRUCK_", "BUS_", "TRAIN_", "TRAMWAY_", "PLANE_", "CARGO_", "BOAT_")
    INFRA_PREFIXES = ("DHN", "GRID", "HVAC_", "HVDC_", "GAS_PIPELINE", "GAS_SUBSEA", "H2_RETRO", "H2_NEW", "H2_SUBSEA", "EFFICIENCY")
    SYNFUEL_PREFIXES = ("SYN_", "BIOMASS_TO_", "BIOWASTE_TO_", "POWER_TO_", "H2_TO_", "OIL_TO_", "GAS_TO_", "METHANOL_TO_", "METHANE_TO_", "DIESEL_TO_", "AMMONIA_TO_", "HABER_", "H2_ELECTRO", "H2_NG", "H2_BIOMASS", "BIOMETHAN")
    STORAGE_PREFIXES = ("TS_", "BATT_", "BEV_BATT", "PHEV_BATT", "CAES", "PHS", "DAM_STORAGE", "GAS_STORAGE", "H2_STORAGE", "DIESEL_STORAGE", "GASOLINE_STORAGE", "LFO_STORAGE", "JET_FUEL_STORAGE", "AMMONIA_STORAGE", "METHANOL_STORAGE", "PT_STORAGE", "ST_STORAGE", "TS_COLD", "CO2_STORAGE")

    for r in tech_rows:
        t = r["Technology"]
        if any(t.startswith(p) for p in TRANSPORT_PREFIXES):
            transport_techs.append(r)
        elif any(t.startswith(p) for p in STORAGE_PREFIXES) or t.endswith("_STORAGE"):
            storage_techs.append(r)
        elif any(t.startswith(p) for p in SYNFUEL_PREFIXES):
            synfuel_techs.append(r)
        elif any(t.startswith(p) for p in INFRA_PREFIXES) or t in ("DHN", "GRID", "EFFICIENCY"):
            infra_techs.append(r)
        else:
            energy_techs.append(r)

    def fmt(v):
        if v is None:
            return ""
        if v == 1e15 or v == 1000000000000000.0:
            return "∞ (uncapped)"
        if v == 100000.0:
            return "100000 (uncapped)"
        return v

    row_num = 6
    def write_section(ws, label, rows_list, start_row):
        r = start_row
        ws.cell(r, 1, label)
        ws.cell(r, 1).font = Font(bold=True, size=11, color="2F5496")
        ws.cell(r, 1).fill = SECTION_FILL
        for c in range(2, 11):
            ws.cell(r, c).fill = SECTION_FILL
        r += 1
        for tr in rows_list:
            vals = [
                tr["Technology"],
                fmt(tr["f_min"]), tr["f_min_source"],
                fmt(tr["f_max"]), tr["f_max_source"],
                fmt(tr["fmin_perc"]), tr["fmin_perc_source"],
                fmt(tr["fmax_perc"]), tr["fmax_perc_source"],
                tr["Rationale"],
            ]
            write_row(ws, r, vals)
            r += 1
        return r

    row_num = write_section(ws_tech, "ENERGY SUPPLY & CONVERSION", energy_techs, row_num)
    row_num = write_section(ws_tech, "TRANSPORT", transport_techs, row_num + 1)
    row_num = write_section(ws_tech, "SYNFUELS & CONVERSION", synfuel_techs, row_num + 1)
    row_num = write_section(ws_tech, "INFRASTRUCTURE", infra_techs, row_num + 1)
    row_num = write_section(ws_tech, "STORAGE", storage_techs, row_num + 1)

    # Freeze pane
    ws_tech.freeze_panes = "A6"

    # ── Sheet 21: Resource Inputs ─────────────────────────────────────────
    if "21_V33_Resource_Inputs" in wb.sheetnames:
        del wb["21_V33_Resource_Inputs"]
    if "21_V37_Resource_Inputs" in wb.sheetnames:
        del wb["21_V37_Resource_Inputs"]
    ws_res = wb.create_sheet("21_V37_Resource_Inputs")

    ws_res.cell(1, 1, "Resource Inputs — v37 Calibration (Finland 2017)")
    ws_res.cell(1, 1).font = Font(bold=True, size=14)
    ws_res.merge_cells("A1:I1")
    ws_res.cell(2, 1, f"Patch chain: {' → '.join(p.replace('.csv','') for p in PATCH_CHAIN)}")
    ws_res.cell(2, 1).font = Font(italic=True, size=9)
    ws_res.merge_cells("A2:I2")

    res_headers = [
        "Resource", "avail_local (GWh)", "Patch (avail_local)",
        "avail_exterior (GWh)", "Patch (avail_ext)",
        "c_op_local (MEUR/GWh)", "Patch (c_op)",
        "Source", "Notes/Hypothesis"
    ]
    res_widths = [22, 18, 22, 20, 22, 20, 22, 35, 55]
    write_header(ws_res, 4, res_headers, res_widths)

    r = 5
    for rr in res_rows:
        vals = [
            rr["Resource"],
            fmt(rr["avail_local"]), rr["avail_local_source"],
            fmt(rr["avail_exterior"]), rr["avail_exterior_source"],
            rr["c_op_local"], rr["c_op_local_source"],
            rr["Source"], rr["Notes"],
        ]
        write_row(ws_res, r, vals)
        r += 1

    # Add local-only resources from base file that may not be in patches
    ws_res.cell(r + 1, 1, "NOTE: Additional local resources from base Resources.csv (not modified by patches)")
    ws_res.cell(r + 1, 1).font = Font(italic=True, color="888888")
    r += 2
    base_local = {
        "WASTE": (11095.02, 0.006079, "Statistics Finland + Eurostat", "Municipal solid waste energy content"),
    }
    for res_name, (avail, cop, src, notes) in base_local.items():
        if res_name not in [rr["Resource"] for rr in res_rows]:
            write_row(ws_res, r, [res_name, avail, "base", "", "", cop, "base", src, notes])
            r += 1

    # Add RES_HYDRO
    ws_res.cell(r + 1, 1, "Resource availability limits (non-fuel)")
    ws_res.cell(r + 1, 1).font = Font(bold=True, size=11, color="2F5496")
    ws_res.cell(r + 1, 1).fill = SECTION_FILL
    for c in range(2, 10):
        ws_res.cell(r + 1, c).fill = SECTION_FILL
    r += 2
    write_row(ws_res, r, [
        "RES_HYDRO", 14600, "v24_hydro_cap", "", "",
        "", "", "Statistics Finland + Fingrid",
        "Hydro resource capped at 14,600 GWh to match 2017 production"
    ])

    ws_res.freeze_panes = "A5"

    # ── Update 12_Validation_2017 ─────────────────────────────────────────
    if "12_Validation_2017" in wb.sheetnames:
        del wb["12_Validation_2017"]
    ws_val = wb.create_sheet("12_Validation_2017")

    ws_val.cell(1, 1, "Validation Scorecard — v37 (Finland 2017)")
    ws_val.cell(1, 1).font = Font(bold=True, size=14)
    ws_val.merge_cells("A1:H1")
    ws_val.cell(2, 1, "Run: 20260323_173930__v37_solar | Score: 1.2% (14 scored metrics)")
    ws_val.cell(2, 1).font = Font(italic=True, size=10)
    ws_val.merge_cells("A2:H2")

    val_headers = ["Metric", "Model (TWh)", "Target (TWh)", "Error (%)", "Weight", "Status", "Structural Note", "Source"]
    val_widths = [22, 14, 14, 12, 10, 12, 45, 40]
    write_header(ws_val, 4, val_headers, val_widths)

    scorecard = [
        ("CO2", 41.25, 41.2, "0.1%", 2.0, "Scored", "", "Statistics Finland GHG inventory"),
        ("PE_GAS", 20.00, 20.0, "0.0%", 1.0, "Scored", "Binding cap (v22)", "Statistics Finland Energy Balance"),
        ("PE_COAL", 35.00, 35.0, "0.0%", 1.5, "Scored", "Binding cap (v16)", "Statistics Finland Energy Balance"),
        ("PE_NUCLEAR", 65.08, 65.0, "0.1%", 1.0, "Scored", "", "IAEA PRIS + TVO/Fortum"),
        ("PE_HYDRO", 14.60, 15.0, "2.7%", 0.8, "Scored", "Structural (hydro cap v24)", "Fingrid + Statistics Finland"),
        ("PE_WIND", 4.79, 5.0, "4.1%", 0.8, "Scored", "Wind capacity structural", "Finnish Wind Power Association"),
        ("PE_OIL", 80.87, 82.0, "1.4%", 1.5, "Scored", "Oil floor (fmin_perc=0.10)", "Statistics Finland Energy Balance"),
        ("PE_BIOMASS", 96.80, 100.0, "3.2%", 1.5, "Scored", "Oil-biomass tradeoff", "Statistics Finland + Luke"),
        ("ELEC_NUCLEAR", 20.82, 21.6, "3.6%", 1.5, "Scored", "Nuclear capacity ceiling", "Statistics Finland Electricity"),
        ("ELEC_HYDRO", 14.60, 14.6, "0.0%", 1.2, "Scored", "", "Statistics Finland Electricity"),
        ("ELEC_WIND", 4.79, 4.8, "0.1%", 1.0, "Scored", "", "Statistics Finland Electricity"),
        ("ELEC_CHP", 20.49, 20.735, "1.2%", 1.2, "Scored", "", "Statistics Finland CHP data"),
        ("ELEC_CONDENSATION", 3.29, 3.284, "0.3%", 0.8, "Scored", "", "Statistics Finland Electricity"),
        ("ELEC_SOLAR", 0.044, 0.044, "0.7%", 0.3, "Scored", "Small absolute value", "Finnish Energy Authority"),
        ("ELEC_GAS", 4.37, 3.2, "36.7%", 0.0, "Info", "CHP co-production + binding gas cap — structural", "Statistics Finland Electricity"),
        ("ELEC_IMPORTS", 19.54, 20.426, "4.4%", 0.0, "Info", "Balance residual", "Fingrid + Statistics Finland"),
        ("HEAT_DHN", 52.91, 36.5, "45.0%", 0.0, "Info", "DHN share applied to all sectors incl. industry — structural", "Finnish Energy / Statistics Finland"),
    ]

    for i, row_data in enumerate(scorecard):
        r = 5 + i
        write_row(ws_val, r, list(row_data))
        if row_data[5] == "Info":
            for c in range(1, 9):
                ws_val.cell(r, c).font = Font(italic=True, color="888888")

    ws_val.freeze_panes = "A5"

    # ── Update 18_Change_Log ──────────────────────────────────────────────
    ws_log = wb["18_Change_Log"]
    next_row = ws_log.max_row + 1
    log_entries = [
        (date.today().isoformat(), "Copilot+User",
         "Fixed grid losses: Misc_indep.json loss_network.ELECTRICITY 0.0864→0.05",
         "Data/2017/00_INDEP/Misc_indep.json",
         "Finnish ~5% actual grid losses; EU-avg 8.64% caused ~3.1 TWh PE overestimate; v35 regression addressed in v36"),
        (date.today().isoformat(), "Copilot+User",
         "v36: DEC_BOILER_WOOD.fmin_perc=0.20, DHN_COGEN_GAS.f_min=1.42 GW",
         "calibration/patches/v36_grid_losses_recalib.csv",
         "Recalibration after grid-loss fix; score 1.4%; 17 patches total"),
        (date.today().isoformat(), "Copilot+User",
         "v37: PV_ROOFTOP.f_min=0.055 GW",
         "calibration/patches/v37_solar.csv",
         "ELEC_SOLAR 8.5%→0.7%; score 1.2%, new best; 18 patches total"),
        (date.today().isoformat(), "Copilot+User",
         "Updated 12_Validation_2017 to v37 scorecard",
         "Finland_MASTER_Calibration_old_UPDATED.xlsx",
         "v37 score=1.2%, 14 scored + 3 informational metrics"),
        (date.today().isoformat(), "Copilot+User",
         "Added 20_V37_Tech_Constraints sheet",
         "Finland_MASTER_Calibration_old_UPDATED.xlsx",
         "All fmin/fmax/fperc from v37 patch chain (18 patches) with sources and rationale"),
        (date.today().isoformat(), "Copilot+User",
         "Added 21_V37_Resource_Inputs sheet",
         "Finland_MASTER_Calibration_old_UPDATED.xlsx",
         "All resource avail/prices from v37 patch chain with sources"),
    ]
    for entry in log_entries:
        for c, v in enumerate(entry, 1):
            ws_log.cell(next_row, c, v)
        next_row += 1

    # ── Save ──────────────────────────────────────────────────────────────
    wb.save(EXCEL)
    print(f"Saved to {EXCEL}")
    print(f"  - 20_V37_Tech_Constraints: {len(tech_rows)} technologies")
    print(f"  - 21_V37_Resource_Inputs: {len(res_rows)} resources")
    print(f"  - 12_Validation_2017: {len(scorecard)} metrics (v37, score=1.2%)")

if __name__ == "__main__":
    main()
