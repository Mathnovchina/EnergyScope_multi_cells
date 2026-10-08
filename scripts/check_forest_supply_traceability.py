"""Read-only traceability check of the Finnish wood supply steps (WOOD_FI1..FI5).

Recomputes the 2035 scenario availabilities from ENSPRESO, compares them with the
scenario CSVs, the patch files, the values actually applied in the 9 canonical runs
(paper1/results) and the GHG factors implied by the run outputs.
Nothing is written. Run from the repository root:  python scripts/check_forest_supply_traceability.py
"""
import csv
import glob
import json
import os
import sys

import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "Data", "exogenous_data", "ENSPRESO", "ENSPRESO_BIOMASS.xlsx")
DATA = os.path.join(ROOT, "Data", "2035")
RESULTS = os.path.join(ROOT, "paper1", "results")
PATCHES = os.path.join(ROOT, "calibration", "patches")

PJ_TO_GWH = 1000 / 3.6
STEPS = {
    "WOOD_FI1": ["MINBIOWOOa"],
    "WOOD_FI2": ["MINBIOFRSR1"],
    "WOOD_FI3": ["MINBIOWOOW1", "MINBIOWOOW1a"],
    "WOOD_FI4": ["MINBIOWOO", "MINBIOFRSR1a"],
}
fails = []


def check(label, got, want, tol=1.0):
    ok = abs(float(got) - float(want)) <= tol
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: {got:.5g} vs {want:.5g}")
    if not ok:
        fails.append(label)


def stat_fi_wood_fuels(years=("2015", "2017", "2019", "2022", "2023", "2024")):
    """Wood fuel consumption incl. black liquor from the open Statistics Finland API (table 12vq)."""
    import urllib.request

    url = "https://pxdata.stat.fi/PxWeb/api/v1/en/StatFin/ehk/12vq.px"
    hdr = {"User-Agent": "Mozilla/5.0", "Content-Type": "application/json"}
    query = {"query": [
        {"code": "timeperiod_y", "selection": {"filter": "item", "values": list(years)}},
        {"code": "energia_36_20201012", "selection": {"filter": "item", "values": ["1.3", "1.3.1", "1.3.2", "1.3.3"]}},
        {"code": "contentscode", "selection": {"filter": "item", "values": ["maara_gwh"]}},
    ], "response": {"format": "json-stat2"}}
    res = json.loads(urllib.request.urlopen(
        urllib.request.Request(url, data=json.dumps(query).encode(), headers=hdr), timeout=60).read())
    labels = [list(res["dimension"][i]["category"]["label"].values()) for i in res["id"]]
    vals, k = res["value"], 0
    for y in labels[0]:
        row = {}
        for src in labels[1]:
            row[src] = vals[k]
            k += 1
        print(f"  [INFO] {y}: " + " | ".join(f"{s[:34]} = {v}" for s, v in row.items()))


def load_enspreso():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    energy, cost = {}, {}
    for r in wb["ENER - NUTS0 EnergyCom"].iter_rows(min_row=2, values_only=True):
        if r[2] == "FI":
            energy[(r[1], r[0], r[3])] = r[4] * PJ_TO_GWH
    for r in wb["COST - NUTS0 EnergyCom"].iter_rows(min_row=2, values_only=True):
        if r[2] == "FI":
            cost[(r[1], r[0], r[3])] = r[4] * 3.6  # EUR2010/GJ -> EUR2010/MWh
    return energy, cost


def step_value(energy, scen, year, step):
    return sum(energy[(scen, year, c)] for c in STEPS[step])


def read_resources_csv(path):
    df = pd.read_csv(path, index_col=0)
    return df


def main():
    energy, cost = load_enspreso()

    print("\n1. ENSPRESO -> scenario availability (GWh/yr) vs scenario CSVs")
    csvs = {s: read_resources_csv(os.path.join(DATA, "FI", f"Resources_{s}.csv"))
            for s in ("S1_BES", "S2_NFS", "S3_BDS")}
    for step in STEPS:
        mid = (step_value(energy, "ENS_Med", 2030, step) + step_value(energy, "ENS_Med", 2040, step)) / 2
        check(f"S2 {step} = mean(ENS_Med 2030, 2040)", csvs["S2_NFS"].loc[step, "avail_local"], mid)
        check(f"S1 {step} = S2 x 1.2", csvs["S1_BES"].loc[step, "avail_local"], mid * 1.2)
        raw = step_value(energy, "ENS_Low", 2030, step)
        print(f"  [INFO] S3 {step}: raw ENS_Low 2030 = {raw:.0f}; final (hand-revised) = "
              f"{csvs['S3_BDS'].loc[step, 'avail_local']:.0f}")

    print("\n2. Patches vs scenario CSVs")
    for scen, patch in (("S1_BES", "fi_forest_S1_BES.csv"), ("S3_BDS", "fi_forest_S3_BDS.csv")):
        with open(os.path.join(PATCHES, patch), newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                res, par, val = row["technology_or_resource"], row["parameter"], float(row["value"])
                check(f"{scen} patch {res} {par}", val, csvs[scen].loc[res, par], tol=1e-9)

    print("\n3. Baseline Resources.csv equals S2 for wood steps")
    base = read_resources_csv(os.path.join(DATA, "FI", "Resources.csv"))
    for step in list(STEPS) + ["WOOD_FI5"]:
        check(f"baseline {step} avail", base.loc[step, "avail_local"], csvs["S2_NFS"].loc[step, "avail_local"], 1e-9)

    print("\n4. Values applied in the 9 canonical runs (paper1/results)")
    for d in sorted(glob.glob(os.path.join(RESULTS, "*"))):
        name = os.path.basename(d)
        scen = name.split("__")[0]
        res = pd.read_csv(os.path.join(d, "Resources.csv")).set_index("Resources")
        gwp = pd.read_csv(os.path.join(d, "Gwp_breakdown.csv")).set_index("Elements")
        for step in list(STEPS):
            check(f"{name} {step} avail", res.loc[step, "avail_local"], csvs[scen].loc[step, "avail_local"], 1e-6)
        used = sum(res.loc[s, "R_year_local"] for s in STEPS)
        print(f"  [INFO] {name}: domestic wood used = {used / 1000:.1f} TWh; "
              f"BIOMASS_RESIDUES avail in run = {res.loc['BIOMASS_RESIDUES', 'avail_local']:.0f} "
              f"(scenario CSV: {csvs[scen].loc['BIOMASS_RESIDUES', 'avail_local']:.0f})")

    print("\n5. GHG factors implied by run outputs vs 02_REF_REGION/Resources.csv (tCO2/GWh)")
    ref = {}
    with open(os.path.join(DATA, "02_REF_REGION", "Resources.csv"), newline="", encoding="utf-8") as f:
        for row in csv.reader(f):
            if len(row) > 5 and row[2].startswith("WOOD"):
                ref[row[2]] = float(row[5])
    print("  [INFO] ref values (ktCO2/GWh):", {k: v for k, v in ref.items()})
    verified = set()
    for d in sorted(glob.glob(os.path.join(RESULTS, "*"))):
        res = pd.read_csv(os.path.join(d, "Resources.csv")).set_index("Resources")
        gwp = pd.read_csv(os.path.join(d, "Gwp_breakdown.csv")).set_index("Elements")
        for step in list(STEPS) + ["WOOD_FI5"]:
            q = res.loc[step, "R_year_local"]
            if q > 100 and step not in verified:
                check(f"{os.path.basename(d)} {step} implied gwp_op", gwp.loc[step, "GWP_op"] / q * 1000, ref[step] * 1000, 0.05)
                verified.add(step)
    print("  [INFO] not verifiable from run outputs (never used):", sorted(set(list(STEPS) + ["WOOD_FI5"]) - verified))

    print("\n6. ENSPRESO cost cross-check (EUR2010/MWh) vs model c_op (EUR/MWh)")
    for step, codes in STEPS.items():
        for c in codes:
            m = (cost[("ENS_Med", 2030, c)] + cost[("ENS_Med", 2040, c)]) / 2
            print(f"  [INFO] {step} {c}: ENSPRESO Med 2035 = {m:.1f}; Low 2030 = {cost[('ENS_Low', 2030, c)]:.1f}; "
                  f"model S2 = {csvs['S2_NFS'].loc[step, 'c_op_local'] * 1000:.1f}")

    print("\n7. Moenkkoenen -> S3 residue chain (Mm3 -> TWh)")
    max_sust = 78.2 / 0.96
    ceiling = max_sust * 0.60
    print(f"  [INFO] max sustainable {max_sust:.1f} Mm3; 60% ceiling {ceiling:.1f} Mm3")
    for frac in (0.15, 0.175, 0.20):
        for eff in (0.30, 0.40):
            vol = ceiling * frac * eff
            print(f"  [INFO] residue frac {frac:.3f}, collection {eff:.2f}: {vol:.2f} Mm3 -> "
                  f"{vol * 1.5:.2f} TWh at 1.5 MWh/m3 | {vol * 1.96:.2f} TWh at 1.96 MWh/m3")
    print(f"  [INFO] LUKE 2.6 Mm3 = 5.1 TWh implies {5.1 / 2.6:.2f} MWh/m3 (not 1.5); 10.1 Mm3 = 20 TWh implies {20 / 10.1:.2f}")

    print("\n8. T1 evidence: biomass commodities behind each forestry energy commodity (FI, ENS_Med 2030, GWh)")
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    comp = {}
    allb = set()
    for r in wb["ENER - NUTS2 BioCom E"].iter_rows(min_row=2, values_only=True):
        allb.add(r[5])
        if r[2] == "FI" and r[0] == 2030 and r[1] == "ENS_Med":
            comp[(r[4], r[5])] = comp.get((r[4], r[5]), 0) + r[6] * PJ_TO_GWH
    for (ecom, bcom), v in sorted(comp.items()):
        if ecom.startswith(("MINBIOWOO", "MINBIOFRSR")):
            print(f"  [INFO] {ecom:13s} <- {bcom:15s} {v:8.0f}")
    hits = sorted(b for b in allb if b and any(k in str(b).lower() for k in ("liquor", "bark", "pulp", "black")))
    print(f"  [INFO] {len(allb)} bio-commodities in workbook; names containing liquor/bark/pulp/black: {hits or 'none'}")
    check("MINBIOWOOa NUTS2 sum (C&P_RW) vs NUTS0", comp[("MINBIOWOOa", "C&P_RW")],
          step_value(energy, "ENS_Med", 2030, "WOOD_FI1"), tol=20)

    if "--online" in sys.argv:
        print("\n9. T1 magnitude test (needs internet): Statistics Finland table 12vq, GWh")
        stat_fi_wood_fuels()

    print("\nRESULT:", "ALL PASS" if not fails else f"{len(fails)} FAIL: {fails}")


if __name__ == "__main__":
    main()
