"""Generate the meeting-opener deck for the IIASA G4M coupling discussion.
Style echoes Docs/Finland_forest_bordenave (1).pptx (16:9, Arial, teal/green accents).
Output: Docs/Finland_forest_IIASA_meeting_2026-10-09.pptx
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = r"c:\Users\borde\OneDrive\Bureau\model\EnergyScope_multi_cells"
PLOTS = os.path.join(ROOT, "plots")
VAL = os.path.join(PLOTS, "validation_2017")
OUT = os.path.join(ROOT, "Docs", "Finland_forest_IIASA_meeting_2026-10-09.pptx")

# palette (nods to old deck dk2 teal-green + scenario colours)
INK = RGBColor(0x1F, 0x2D, 0x2D)
TEAL = RGBColor(0x15, 0x81, 0x58)       # dk2 of old deck
TEAL_D = RGBColor(0x0E, 0x5A, 0x3E)
GREY = RGBColor(0x5A, 0x63, 0x63)
LIGHT = RGBColor(0xF3, 0xF6, 0xF4)
BAND = RGBColor(0xE8, 0xF0, 0xEB)
S1 = RGBColor(0xC0, 0x39, 0x2B)         # BES red
S2 = RGBColor(0x7F, 0x8C, 0x8D)         # NFS grey
S3 = RGBColor(0x27, 0xAE, 0x60)         # BDS green
AMBER = RGBColor(0xB9, 0x77, 0x0E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def box(s, l, t, w, h):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def rect(s, l, t, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp


def set_runs(para, parts, size, color=INK, bold=False, italic=False):
    para.clear() if False else None
    first = True
    for seg in parts:
        if isinstance(seg, tuple):
            txt, kw = seg
        else:
            txt, kw = seg, {}
        r = para.add_run()
        r.text = txt
        r.font.name = FONT
        r.font.size = Pt(kw.get("size", size))
        r.font.bold = kw.get("bold", bold)
        r.font.italic = kw.get("italic", italic)
        r.font.color.rgb = kw.get("color", color)
    return para


def title_banner(s, kicker, title):
    rect(s, 0, 0, 13.333/1, 1.15, TEAL)
    rect(s, 0, 1.15, 13.333/1, 0.06, RGBColor(0x50, 0xB4, 0x32))
    _, tf = box(s, 0.55, 0.12, 12.2, 1.0)
    p = tf.paragraphs[0]
    set_runs(p, [(kicker.upper(), {"size": 11, "color": RGBColor(0xC9, 0xE8, 0xD6), "bold": True})], 11)
    p2 = tf.add_paragraph()
    set_runs(p2, [(title, {"size": 24, "color": WHITE, "bold": True})], 24)


def footer(s, n):
    _, tf = box(s, 0.5, 7.03, 9, 0.35)
    p = tf.paragraphs[0]
    set_runs(p, [("M. Bordenave · Forest–Energy coupling for Finland · IIASA meeting, 9 Oct 2026", {"size": 8.5, "color": GREY})], 8.5)
    _, tf2 = box(s, 12.4, 7.03, 0.6, 0.35)
    p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.RIGHT
    set_runs(p2, [(str(n), {"size": 9, "color": GREY, "bold": True})], 9)


def bullets(s, items, l, t, w, h, size=14, gap=6):
    _, tf = box(s, l, t, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        lvl = it[0] if isinstance(it, tuple) and isinstance(it[0], int) else 0
        content = it[1] if (isinstance(it, tuple) and isinstance(it[0], int)) else it
        p.level = lvl
        marker = "—  " if lvl == 0 else "·   "
        col = TEAL_D if lvl == 0 else GREY
        runs = content if isinstance(content, list) else [content]
        set_runs(p, [(marker, {"color": col, "bold": True, "size": size})] + runs, size)
    return tf


N = [0]
def nxt():
    N[0] += 1
    return N[0]


# ---------- 1 TITLE ----------
s = slide()
rect(s, 0, 0, 13.333, 7.5, TEAL)
rect(s, 0, 4.55, 13.333, 0.05, RGBColor(0x50, 0xB4, 0x32))
_, tf = box(s, 0.8, 1.5, 11.7, 2.6)
p = tf.paragraphs[0]
set_runs(p, [("From Bioeconomy Promise to Biodiversity Ceiling", {"size": 34, "color": WHITE, "bold": True})], 34)
p2 = tf.add_paragraph(); p2.space_before = Pt(6)
set_runs(p2, [("The role of forest biomass in Finnish decarbonisation —", {"size": 20, "color": RGBColor(0xD6, 0xEC, 0xDE)})], 20)
p3 = tf.add_paragraph()
set_runs(p3, [("and the next phase: coupling G4M, disturbances and the forest carbon sink", {"size": 20, "color": RGBColor(0xD6, 0xEC, 0xDE), "italic": True})], 20)
_, tf = box(s, 0.8, 4.8, 11.7, 1.6)
p = tf.paragraphs[0]
set_runs(p, [("Matthieu Bordenave", {"size": 18, "color": WHITE, "bold": True})], 18)
p2 = tf.add_paragraph()
set_runs(p2, [("University of Pisa · IIASA · UCLouvain", {"size": 13, "color": RGBColor(0xC9, 0xE8, 0xD6)})], 13)
p3 = tf.add_paragraph(); p3.space_before = Pt(8)
set_runs(p3, [("Working meeting with IIASA — 9 October 2026", {"size": 13, "color": RGBColor(0xC9, 0xE8, 0xD6), "italic": True})], 13)

# ---------- 2 AGENDA ----------
s = slide(); title_banner(s, "Roadmap", "What we will cover today")
items = [
    [("1. Context & motivation", {"bold": True}), ("  —  forest biomass as Finland's main domestic decarbonisation lever", {"color": GREY})],
    [("2. Paper 1 (near-complete)", {"bold": True}), ("  —  method, stepwise wood supply curves, 3×3 scenario matrix, headline results", {"color": GREY})],
    [("3. Two methodological gaps", {"bold": True}), ("  —  static curves · no forest carbon sink", {"color": GREY})],
    [("4. Research questions for the next phase", {"bold": True})],
    [("5. Proposal to IIASA", {"bold": True}), ("  —  G4M supply curves · disturbance shocks · sink coupling · European/BeWhere extension", {"color": GREY})],
    [("6. The IIASA model landscape", {"bold": True}), ("  —  G4M, PICUS/PREBAS, FLAM, BeWhere, ForestNavigator D3.4", {"color": GREY})],
    [("7. Path forward & open questions", {"bold": True}), ("  —  what we need from the G4M / ForestNavigator teams", {"color": GREY})],
]
bullets(s, items, 0.7, 1.55, 12.0, 5.2, size=15, gap=10)
footer(s, nxt())

# ---------- 3 CONTEXT ----------
s = slide(); title_banner(s, "Context & motivation", "Why forest biomass, why supply curves")
bullets(s, [
    [("Forest biomass supplies ~25% of Finland's primary energy", {"bold": True}), ("  — the main domestic lever for deep decarbonisation.", {"color": GREY})],
    [("But availability is ", {}), ("not a single fixed number", {"bold": True}), (": it depends on forest management, competing industrial uses (pulp, paper, sawnwood), ecological limits, and — increasingly — natural disturbances.", {})],
    [("Colla et al. (2022)", {"bold": True}), (" showed for Belgium that representing wood as a ", {}), ("multi-step marginal-cost supply curve", {"bold": True}), (" materially changes its optimal allocation under deep decarbonisation.", {})],
    [("Paper 1 applies and extends this to Finland", {"bold": True}), (" — a structurally richer, more policy-contested case (bioeconomy vs biodiversity vs climate).", {})],
    [("Next step:", {"bold": True, "color": TEAL_D}), (" make the curves ", {}), ("endogenous to forest dynamics", {"bold": True}), (", so management intensity and disturbance shocks propagate into the energy system through both wood availability and the forest carbon balance.", {})],
], 0.7, 1.55, 12.1, 5.2, size=15, gap=12)
footer(s, nxt())

# ---------- 4 TOOL ----------
s = slide(); title_banner(s, "The tool", "EnergyScope — whole-energy-system optimisation")
bullets(s, [
    [("EnergyScope TD", {"bold": True}), (" (Limpens et al. 2019): open-source LP that co-optimises investment and hourly operation of the whole energy system — electricity, heat, mobility, non-energy demand — to minimise total annual cost under a GHG constraint.", {})],
    [("Typical-days representation", {"bold": True}), (" keeps the hourly detail tractable (12 typical days here).", {})],
    [("Multi-cell version (ESMC)", {"bold": True}), (" places Finland inside a pan-European grid; this study is the Finnish national deep-dive.", {})],
    [("Biomass enters as resources with a yearly availability cap, a marginal cost and an embedded (supply-chain) GHG factor", {"bold": True}), (" — exactly where the supply-curve representation lives.", {})],
], 0.7, 1.55, 12.1, 3.2, size=15, gap=12)
rect(s, 0.7, 5.3, 12.0, 1.5, BAND)
_, tf = box(s, 0.95, 5.45, 11.6, 1.25)
p = tf.paragraphs[0]
set_runs(p, [("Binding climate constraint (today): ", {"bold": True, "color": TEAL_D}),
             ("Σ_r CO₂_net[r] ≤ gwp_limit", {"bold": True, "italic": True}),
             ("   — energy-system combustion emissions only. Biomass combustion is carbon-neutral by assumption; there is no term for the forest sink.", {"color": INK})], 14)
footer(s, nxt())

# ---------- 5 PAPER 1 OVERVIEW ----------
s = slide(); title_banner(s, "Paper 1 — what is done", "Three building blocks, results are stylised")
# three columns
cols = [
    ("Calibrated 2017 baseline", S2, [
        "Reproduces the national energy balance to ~1.2% weighted error over 14 indicators",
        "CO₂ within 0.1% (41.2 MtCO₂)",
        "Residuals documented: district heat +43%, gas-CHP +37% (aggregation artefacts)",
    ]),
    ("Stepwise wood supply curves (2035)", S3, [
        "Single WOOD resource → 4 domestic steps + import backstop",
        "Dispatched in merit order (cheapest first)",
        "Three forest-management scenarios S1 / S2 / S3",
    ]),
    ("3 × 3 scenario matrix", S1, [
        "{S1, S2, S3} × {unconstrained, −95% GHG, −95% no-nuclear}",
        "9 canonical runs, reproducible (tag paper1-finland-forest-v1.0)",
        "Headline findings → next slides",
    ]),
]
x = 0.7
for name, c, pts in cols:
    rect(s, x, 1.55, 3.95, 0.62, c)
    _, tf = box(s, x + 0.12, 1.63, 3.7, 0.5)
    set_runs(tf.paragraphs[0], [(name, {"size": 13.5, "color": WHITE, "bold": True})], 13.5)
    rect(s, x, 2.17, 3.95, 4.1, LIGHT)
    _, tf = box(s, x + 0.14, 2.3, 3.67, 3.9)
    for i, pt in enumerate(pts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        set_runs(p, [("•  ", {"color": c, "bold": True, "size": 12.5}), (pt, {"size": 12.5})], 12.5)
    x += 4.18
footer(s, nxt())

# ---------- 6 SUPPLY CURVES ----------
s = slide(); title_banner(s, "Paper 1 — supply curves", "Stepwise Finnish wood supply, by management scenario")
if os.path.exists(os.path.join(PLOTS, "biomass_supply_curves_fi_2035.png")):
    s.shapes.add_picture(os.path.join(PLOTS, "biomass_supply_curves_fi_2035.png"), Inches(0.5), Inches(1.5), height=Inches(3.9))
# right-side legend of steps
_, tf = box(s, 8.3, 1.5, 4.7, 4.2)
rows = [
    ("FI1", "Industrial by-products — black liquor, bark, sawdust", "11"),
    ("FI2", "Logging residues (branches, tops)", "22 (26 S3)"),
    ("FI3", "Secondary woodchips / sawdust", "27"),
    ("FI4", "Fuelwood + landscape-care wood", "33"),
    ("FI5", "Import backstop (Baltic / Nordic)", "70"),
]
p = tf.paragraphs[0]
set_runs(p, [("Five merit-order steps  (€/MWh)", {"size": 13, "bold": True, "color": TEAL_D})], 13)
for code, desc, cost in rows:
    pp = tf.add_paragraph(); pp.space_after = Pt(5)
    set_runs(pp, [(f"{code}  ", {"bold": True, "size": 12, "color": INK}), (f"{desc}  ", {"size": 11.5}), (f"— {cost}", {"size": 11.5, "bold": True, "color": GREY})], 12)
_, tf = box(s, 0.5, 5.55, 12.3, 1.35)
p = tf.paragraphs[0]
set_runs(p, [("Domestic ceilings: ", {"bold": True}),
             ("S1-BES ~122 TWh", {"bold": True, "color": S1}), (" · ", {}),
             ("S2-NFS ~102 TWh", {"bold": True, "color": S2}), (" · ", {}),
             ("S3-BDS 40 TWh", {"bold": True, "color": S3}),
             ("   (S1 +20% and S3 40 TWh are stylised/revised, not raw ENSPRESO).", {"italic": True, "color": GREY})], 13)
p2 = tf.add_paragraph(); p2.space_before = Pt(4)
set_runs(p2, [("Note: ", {"bold": True, "color": AMBER}),
              ("FI1 proxies Finland's captive by-product stream, dominated by black liquor (~43 TWh, 2017) — absent from ENSPRESO; flagged for IIASA.", {"italic": True, "color": GREY, "size": 11.5})], 11.5)
footer(s, nxt())

# ---------- 7 RESULTS TABLE ----------
s = slide(); title_banner(s, "Paper 1 — results", "The 2035 scenario × GHG matrix")
data = [
    ["Scenario", "GHG case", "Cost\n(bn€/y)", "CO₂\n(Mt/y)", "Reduction\nvs 2017", "Wood used\n(TWh)"],
    ["S1 BES", "Unconstrained", "22.40", "13.02", "68%", "54.5"],
    ["S1 BES", "−95% (nuclear)", "22.92", "2.06", "95%", "74.2"],
    ["S1 BES", "−95% (no nuke)", "22.94", "2.06", "95%", "100.4"],
    ["S2 NFS", "Unconstrained", "22.46", "12.57", "70%", "45.4"],
    ["S2 NFS", "−95% (nuclear)", "23.02", "2.06", "95%", "74.2"],
    ["S2 NFS", "−95% (no nuke)", "23.12", "2.06", "95%", "96.2"],
    ["S3 BDS", "Unconstrained", "22.62", "13.10", "68%", "32.0"],
    ["S3 BDS", "−95% (nuclear)", "23.56", "2.06", "95%", "40.0 ✦"],
    ["S3 BDS", "−95% (no nuke)", "23.90", "2.06", "95%", "40.0 ✦"],
]
rowcol = {"S1 BES": S1, "S2 NFS": S2, "S3 BDS": S3}
tbl_shape = s.shapes.add_table(len(data), 6, Inches(0.55), Inches(1.5), Inches(7.6), Inches(5.1))
table = tbl_shape.table
widths = [1.25, 1.75, 1.05, 0.95, 1.2, 1.4]
for j, wv in enumerate(widths):
    table.columns[j].width = Inches(wv)
for i, row in enumerate(data):
    for j, val in enumerate(row):
        cell = table.cell(i, j)
        cell.margin_top = Pt(2); cell.margin_bottom = Pt(2); cell.margin_left = Pt(5); cell.margin_right = Pt(3)
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if j > 1 else PP_ALIGN.LEFT
        r = p.add_run(); r.text = val; r.font.name = FONT
        if i == 0:
            r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = WHITE
            cell.fill.solid(); cell.fill.fore_color.rgb = TEAL
        else:
            r.font.size = Pt(11)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            if j == 0:
                r.font.bold = True; r.font.color.rgb = rowcol[val]
# messages panel
rect(s, 8.4, 1.5, 4.4, 5.1, BAND)
_, tf = box(s, 8.6, 1.65, 4.05, 4.85)
p = tf.paragraphs[0]
set_runs(p, [("Three headline messages", {"size": 14, "bold": True, "color": TEAL_D})], 14)
msgs = [
    "2035 is already ~70% below 2017 even unconstrained — a structural floor (coal phase-out + brownfield renewables).",
    "Under −95%, S3 (conservation) saturates its 40 TWh ceiling and becomes the most expensive configuration.",
    "Removing nuclear sharply raises the value of biomass in S1/S2 (up to ~100 TWh) — but S3 physically cannot expand (✦ = ceiling binding).",
]
for i, m in enumerate(msgs):
    pp = tf.add_paragraph(); pp.space_before = Pt(8)
    set_runs(pp, [(f"{i+1}.  ", {"bold": True, "color": TEAL_D, "size": 12.5}), (m, {"size": 12.5})], 12.5)
footer(s, nxt())

# ---------- 8 GAPS ----------
s = slide(); title_banner(s, "Two methodological gaps", "What Paper 1 cannot yet represent")
rect(s, 0.7, 1.6, 5.9, 4.9, LIGHT)
rect(s, 0.7, 1.6, 5.9, 0.62, AMBER)
_, tf = box(s, 0.9, 1.68, 5.5, 0.5)
set_runs(tf.paragraphs[0], [("1 · Static & exogenous curves", {"size": 15, "bold": True, "color": WHITE})], 15)
bullets(s, [
    "Curves are read from ENSPRESO potentials and published papers, not from a forest model.",
    "Scenarios differ only in availability and cost — no forest growth dynamics, age structure or disturbances.",
    [("S1 (+20%) and S3 (40 TWh) are stylised / revised", {"bold": True}), (" — the cross-scenario volume differences rest largely on assumptions.", {})],
], 0.95, 2.35, 5.45, 3.9, size=13, gap=10)
rect(s, 6.75, 1.6, 5.9, 4.9, LIGHT)
rect(s, 6.75, 1.6, 5.9, 0.62, S1)
_, tf = box(s, 6.95, 1.68, 5.5, 0.5)
set_runs(tf.paragraphs[0], [("2 · The forest carbon sink is absent", {"size": 15, "bold": True, "color": WHITE})], 15)
bullets(s, [
    "The binding constraint counts energy-system combustion CO₂ only; biomass is carbon-neutral by assumption.",
    "The model sees the S1-vs-S3 climate difference only through supply-chain GHG.",
    [("It cannot represent the dominant effect of harvest intensity", {"bold": True}), (": lower harvest raises standing stock and the sink; intensive harvest depletes it.", {})],
    [("Finnish/EU literature treats this as first-order", {"bold": True}), (" (e.g. Blattert: NFS sink up to ~80, BES ~28 MtCO₂/yr; BDS starts highest then collapses).", {"color": GREY})],
], 7.0, 2.35, 5.45, 3.9, size=13, gap=8)
footer(s, nxt())

# ---------- 9 RESEARCH QUESTIONS ----------
s = slide(); title_banner(s, "Next phase", "Research questions")
rqs = [
    ("RQ1 — Management", "How do contrasting management regimes reshape the dynamic Finnish wood supply curve (quantity, quality, marginal cost, 2025–2060) when derived from G4M rather than a static database?"),
    ("RQ2 — Sink coupling", "When the forest carbon balance (sink/source + HWP) enters the national GHG budget, how does the optimal energy mix and the cost of deep decarbonisation change — and does the S1/S2/S3 climate ranking change?"),
    ("RQ3 — Shocks", "How does a compound disturbance (windstorm → bark-beetle) propagate into the energy system — calamity-wood pulse, multi-year deficit, sink loss — and does it strand biomass-conversion assets?"),
    ("RQ4 — Policy framing", "If the 'neutral' baseline becomes a sink-maximising / LULUCF net-zero scenario, what does economy-wide net-zero (energy + land) imply vs energy-only net-zero?"),
    ("RQ5 — Scaling", "Can the same supply-curve representation scale to Europe and to competing end-uses (construction / HWP) with BeWhere?"),
]
_, tf = box(s, 0.7, 1.5, 12.1, 5.4)
for i, (h, t) in enumerate(rqs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(11)
    set_runs(p, [(h + "  —  ", {"bold": True, "color": TEAL_D, "size": 14}), (t, {"size": 13})], 13)
footer(s, nxt())

# ---------- 10 PROPOSAL A ----------
s = slide(); title_banner(s, "Proposal to IIASA — A", "G4M management-scenario supply curves")
bullets(s, [
    [("Replace the hand-built curves with physically modelled, assortment-level harvest from G4M", {"bold": True}), (" for Finland, 2025–2060, under each regime × climate pathway (GWL2 / GWL3).", {})],
    [("Extract per year:", {"bold": True}), (" harvestable volume by assortment (sawlogs, pulpwood, logging residues, fuelwood, stumps); ecologically-available fractions; net forest + HWP carbon flux.", {})],
    [("Derive the energy-available volume as the residual", {"bold": True}), (" after industrial roundwood demand under each regime — exactly the step we now assume.", {})],
    [("Map assortments → existing WOOD_FI1…FI4 steps", {"bold": True}), (" (volume→energy factors, costs from Finnish statistics / RED II). EnergyScope implementation unchanged.", {})],
    [("Structural question to settle:", {"bold": True, "color": AMBER}), (" FI1 is a captive industrial by-product stream (black liquor, bark, sawdust) driven by processing throughput, not forest-gate harvest — keep from statistics, or bridge from G4M roundwood?", {})],
], 0.7, 1.55, 12.1, 5.2, size=13.5, gap=11)
footer(s, nxt())

# ---------- 11 PROPOSAL B ----------
s = slide(); title_banner(s, "Proposal to IIASA — B", "Disturbance-shock supply curves")
bullets(s, [
    [("Use G4M's disturbance-aware runs (+ ForestNavigator FLAM / bark-beetle / wind) to build time-profiled curves around a compound event", {"bold": True}), (", not a single lower ceiling.", {})],
    [("A disturbance scenario superimposes three effects:", {"bold": True})],
    (1, [("Calamity-wood pulse", {"bold": True}), (" — a large, cheap, time-limited (~18–24 mo) salvage volume, quality-penalised (blue-stain, moisture, LHV −10…−25%).", {})]),
    (1, [("Structural supply deficit", {"bold": True}), (" — reduced availability for 10–15 years from depleted cohorts and damaged area out of rotation.", {})]),
    (1, [("Sink loss", {"bold": True}), (" — an immediate LULUCF emission from killed biomass and foregone growth.", {})]),
    [("Energy-system response studied as a pre-shock → post-shock pair with constrained re-optimisation", {"bold": True}), (" (inherited lower bounds on installed capacity) to expose stranded-asset and transition-delay risk.", {})],
], 0.7, 1.55, 12.1, 5.2, size=13.5, gap=8)
footer(s, nxt())

# ---------- 12 PROPOSAL C ----------
s = slide(); title_banner(s, "Proposal to IIASA — C", "Coupling the forest carbon sink into EnergyScope")
_, tf = box(s, 0.7, 1.5, 12.1, 1.5)
p = tf.paragraphs[0]
set_runs(p, [("Make the national target a land–energy coupled budget — a small, transparent AMPL change:", {"size": 14, "bold": True})], 14)
rect(s, 0.9, 2.25, 11.5, 1.35, BAND)
_, tf = box(s, 1.15, 2.4, 11.0, 1.1)
p = tf.paragraphs[0]
set_runs(p, [("Σ_r CO₂_net[r] ", {"bold": True, "italic": True, "size": 17}),
             ("− S_forest(scenario, year, disturbance)", {"bold": True, "italic": True, "size": 17, "color": TEAL_D}),
             ("  ≤  gwp_limit", {"bold": True, "italic": True, "size": 17})], 17)
p2 = tf.add_paragraph()
set_runs(p2, [("S_forest = exogenous net forest + HWP CO₂ flux from G4M (one national parameter per scenario-year; <0 when a sink).", {"size": 12, "color": GREY, "italic": True})], 12)
bullets(s, [
    [("Harvest intensity and disturbances now act on the energy system through the sink", {"bold": True}), (", not only through wood availability.", {})],
    [("RQ2:", {"bold": True, "color": TEAL_D}), (" does deep decarbonisation become cheaper under conservation (large sink credit), or does the energy-only penalty dominate?", {})],
    [("RQ4 / new S0 scenario:", {"bold": True, "color": TEAL_D}), (" define net-zero over energy + land (Σ_r CO₂_net ≤ S_forest) — a sink-maximising baseline replacing the 'neutral' S2 framing.", {})],
    [("Scope control:", {"bold": True}), (" S_forest is an exogenous input, not endogenised. EnergyScope responds to forest management; it does not optimise it. One-way coupling, tractable for a first paper.", {})],
], 0.7, 3.75, 12.1, 3.0, size=13, gap=9)
footer(s, nxt())

# ---------- 13 PROPOSAL D ----------
s = slide(); title_banner(s, "Proposal to IIASA — D", "European, cross-sectoral extension with BeWhere")
bullets(s, [
    [("Parallel / follow-on track with Fabian Schipfer & Shubham Tiwari (IIASA).", {"bold": True})],
    [("Scale the quality- and cost-differentiated supply-curve representation", {"bold": True}), (" from Finland to Europe, and from energy to competing end-uses (construction / engineered wood / HWP).", {})],
    [("Division of view:", {"bold": True}), (" EnergyScope-Finland provides the national energy-system deep dive; BeWhere provides the spatially explicit, multi-sector, European allocation.", {})],
    [("Shared object:", {"bold": True, "color": TEAL_D}), (" a transferable, disturbance- and management-aware biomass supply-curve interface.", {})],
    [("Open:", {"bold": True, "color": AMBER}), (" is there a Schipfer/Tiwari BeWhere capability for construction / HWP supply chains, and on what timeline for a joint output (paper / proposal)?", {})],
], 0.7, 1.55, 12.1, 5.2, size=14, gap=12)
footer(s, nxt())

# ---------- 14 IIASA MODELS ----------
s = slide(); title_banner(s, "The IIASA model landscape", "What each tool does — and the key caveat for us")
rows = [
    ("G4M", "Geographically explicit economic forest model (NPV). Afforest/deforest/manage from wood prices; harvest, carbon pools, LULUCF flux.", "Assortment split lives in GLOBIOM, not G4M. 0.5° / SimU — Finland resolved as a country but reported EU-27. No deadwood pool found."),
    ("PICUS / PREBAS", "Stand/patch growth models with disturbance modules.", "In ForestNavigator the native PICUS disturbance modules were off; for Northern Europe the growth model is PREBAS — clarify which underlies Finland."),
    ("FLAM", "Process-based daily wildfire model (burned area).", "No Finland application yet (closest: Sweden 2022). Wildfire is secondary for Finland — treat as sensitivity."),
    ("BeWhere", "Spatially explicit MILP for biomass supply chains and plant siting.", "Finland precedent exists (Natarajan 2014). No Schipfer HWP/construction paper located — ask directly."),
    ("ForestNavigator D3.4", "FLAM + G4M + PICUS-derived beetle/wind, 5 arc-min over EU-27 incl. Finland.", "Disturbances modelled as biomass loss only — salvage-wood recovery NOT yet implemented. The ready-made disturbance curve does not exist yet."),
]
tbl_shape = s.shapes.add_table(len(rows) + 1, 3, Inches(0.5), Inches(1.5), Inches(12.33), Inches(5.0))
table = tbl_shape.table
for j, wv in enumerate([2.0, 5.1, 5.23]):
    table.columns[j].width = Inches(wv)
hdr = ["Model", "What it does", "Key caveat for this coupling"]
for j, val in enumerate(hdr):
    cell = table.cell(0, j); cell.fill.solid(); cell.fill.fore_color.rgb = TEAL
    p = cell.text_frame.paragraphs[0]; r = p.add_run(); r.text = val
    r.font.name = FONT; r.font.size = Pt(11.5); r.font.bold = True; r.font.color.rgb = WHITE
for i, (m, d, c) in enumerate(rows, start=1):
    for j, val in enumerate([m, d, c]):
        cell = table.cell(i, j); cell.margin_top = Pt(2); cell.margin_bottom = Pt(2); cell.margin_left = Pt(5)
        cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
        p = cell.text_frame.paragraphs[0]; p.word_wrap = True
        r = p.add_run(); r.text = val; r.font.name = FONT
        r.font.size = Pt(10.5)
        if j == 0:
            r.font.bold = True; r.font.color.rgb = TEAL_D
        if j == 2:
            r.font.color.rgb = RGBColor(0x8A, 0x55, 0x08)
footer(s, nxt())

# ---------- 15 PATH FORWARD ----------
s = slide(); title_banner(s, "Path forward", "Who does what")
steps = [
    ("IIASA — G4M team", "Confirm G4M can deliver for Finland (2025–2060 × regime × GWL2/GWL3): assortment-level harvest, ecologically-available fractions, net forest + HWP flux, disturbance mortality/salvage. Spatial resolution & format."),
    ("M. Bordenave", "Prototype the sink-coupling AMPL term against the existing 2035 matrix using literature sink values (Blattert) as a placeholder — demonstrate the mechanism before G4M data arrive."),
    ("M. Bordenave + IIASA", "Map G4M assortments → WOOD_FI1…FI4; settle the FI1 by-product/black-liquor treatment; adopt best Finnish cost data."),
    ("All", "Agree scenario definitions (incl. the new S0 sink-max / net-zero) and the disturbance event design (compound windstorm → beetle, GWL2 primary)."),
    ("IIASA — Schipfer / Tiwari", "Scope the BeWhere European extension as a follow-on; identify the shared supply-curve interface and a target joint output."),
]
_, tf = box(s, 0.7, 1.5, 12.1, 5.4)
for i, (who, what) in enumerate(steps):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(10)
    set_runs(p, [(f"[{who}]  ", {"bold": True, "color": TEAL_D, "size": 13.5}), (what, {"size": 13})], 13)
footer(s, nxt())

# ---------- 16 OPEN QUESTIONS ----------
s = slide(); title_banner(s, "Open questions for IIASA", "What we need to settle today")
qs = [
    "Does G4M run for Finland at a resolution usable for national supply curves, and output assortment-level harvest + net carbon/HWP flux per management scenario? G4M alone, or via GLOBIOM?",
    "Does G4M output a separate, quantified deadwood carbon pool per scenario-year? (Needed for the sink term and the BDS narrative.)",
    "How do you endogenise the industry-vs-energy split, so energy-available residue is what remains after industrial roundwood demand? What industrial-demand trajectory?",
    "Are ecological availability fractions (deadwood/residue retention, set-aside) applied inside G4M per scenario, or must we post-apply them?",
    "For Finland, is growth/disturbance driven by G4M, PREBAS, or PICUS in your setup?",
    "D3.4 models disturbance as biomass loss only — is a salvage-wood recovery extension available or planned, or do we build the post-processing?",
    "How should calamity-wood carbon be treated in Finland's LULUCF / forest reference level for the coupled GHG budget?",
    "Is there a BeWhere capability for construction / HWP supply chains (Schipfer / Tiwari), and on what timeline?",
]
_, tf = box(s, 0.7, 1.5, 12.1, 5.4)
for i, q in enumerate(qs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(7)
    set_runs(p, [(f"Q{i+1}.  ", {"bold": True, "color": TEAL_D, "size": 13}), (q, {"size": 12.5})], 12.5)
footer(s, nxt())

# ---------- 17 CLOSING ----------
s = slide()
rect(s, 0, 0, 13.333, 7.5, TEAL)
rect(s, 0, 3.6, 13.333, 0.05, RGBColor(0x50, 0xB4, 0x32))
_, tf = box(s, 0.9, 2.2, 11.5, 2.0)
p = tf.paragraphs[0]
set_runs(p, [("From a static biomass ladder …", {"size": 26, "color": RGBColor(0xD6, 0xEC, 0xDE)})], 26)
p2 = tf.add_paragraph(); p2.space_before = Pt(10)
set_runs(p2, [("… to a dynamic, disturbance- and sink-aware forest–energy coupling for Finland", {"size": 26, "color": WHITE, "bold": True})], 26)
_, tf = box(s, 0.9, 5.0, 11.5, 1.2)
p = tf.paragraphs[0]
set_runs(p, [("Discussion — thank you", {"size": 18, "color": RGBColor(0xC9, 0xE8, 0xD6), "italic": True})], 18)

# ---------- 18 BACKUP ----------
s = slide(); title_banner(s, "Backup", "Data provenance & the FI1 / black-liquor flag")
bullets(s, [
    [("Wood volumes:", {"bold": True}), (" ENSPRESO (JRC) commodity codes → 4 steps; S2 = ENS_Med 2035 (interpolated), S1 = S2 ×1.2, S3 = ENS_Low 2030 revised to 40 TWh (LUKE + Mönkkönen audit). All reproduce exactly (1 PJ = 277.78 GWh).", {})],
    [("Costs / GHG:", {"bold": True}), (" model assumptions; GHG from RED II Annex VI via Colla (2022). FI2 = 22 €/MWh matches observed Finnish forest-chip prices.", {})],
    [("FI1 flag (for IIASA):", {"bold": True, "color": AMBER}), (" ENSPRESO formally labels the FI1 commodity as stemwood chips; we re-purpose its volume as a proxy for the captive by-product stream (black liquor dominant, 43 TWh in 2017) because black liquor is absent from ENSPRESO. Cost 11 €/MWh is conservative; flagged for a physically grounded G4M split.", {})],
    [("Reproducibility:", {"bold": True}), (" git tag paper1-finland-forest-v1.0; 9 canonical runs; audit script check_forest_supply_traceability.py (ALL PASS).", {})],
], 0.7, 1.55, 12.1, 5.2, size=13, gap=11)
footer(s, nxt())

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
