
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime

file_path = r'c:\Users\borde\OneDrive\Bureau\model\EnergyScope_multi_cells\Data\exogenous_data\Finland_Calibration_MASTER.xlsx'

def append_df_to_excel(filename, df, sheet_name, startrow=None, replace=False):
    try:
        if replace:
            # Not implemented safe replace in this snippet without losing formatting, 
            # but we are appending.
            pass
        
        book = load_workbook(filename)
        if sheet_name not in book.sheetnames:
            print(f"Sheet {sheet_name} not found.")
            return

        writer = pd.ExcelWriter(filename, engine='openpyxl') 
        writer.book = book
        
        # Access the sheet
        writer.sheets = {ws.title: ws for ws in book.worksheets}
        
        if startrow is None:
            startrow = writer.sheets[sheet_name].max_row

        df.to_excel(writer, sheet_name=sheet_name, startrow=startrow, index=False, header=False)
        writer.close()
        print(f"Appended to {sheet_name}")
        
    except Exception as e:
        print(f"Error appending to {sheet_name}: {e}")

def main():
    print(f"Updating {file_path}")
    today = datetime.now().strftime("%Y-%m-%d")
    
    # 1. Critical Changes
    changes_data = [
        ['High', 'WIND_ONSHORE 2017', '2035 constraint (f_min > 3GW) infeasible for 2017 (<2GW)', 'Update f_min using 2017 specific data', 'Done', 'Critical - Feasibility'],
        ['High', 'NUCLEAR 2035/2050', 'Forced to 0 GW in default config', 'Updated CSVs with realistic capacity (4.36 GW)', 'Done', 'Critical - Base Load']
    ]
    df_changes = pd.DataFrame(changes_data, columns=['Priority', 'Item', 'Issue', 'Fix', 'Status', 'Impact'])
    append_df_to_excel(file_path, df_changes, '3_Critical_Changes')

    # 2. Next Steps
    steps_data = [
        ['Calibration', 1, 'Initial 2017 Model Run', 'Ready', '1 day', 'Data structure complete'],
        ['Post-Processing', 2, 'Validate 2017 Output vs Stats', 'Pending', '1 day', 'Check production mix']
    ]
    df_steps = pd.DataFrame(steps_data, columns=['Phase', 'Priority', 'Task', 'Status', 'Expected_Duration', 'Notes'])
    append_df_to_excel(file_path, df_steps, '11_Next_Steps')

    # 3. Issues Decisions
    issues_data = [
        [today, '2017 demand data source', 'Use Energy in Finland 2022 (via Calibration file)', 'Validated source', 'Closed'],
        [today, 'Cost Learning Curves', 'Use static 2035 costs for now', 'No learning script in repo. Impact considered low for 2017 calib.', 'Open']
    ]
    df_issues = pd.DataFrame(issues_data, columns=['Date', 'Issue', 'Decision', 'Rationale', 'Status'])
    append_df_to_excel(file_path, df_issues, '12_Issues_Decisions')

if __name__ == "__main__":
    main()
